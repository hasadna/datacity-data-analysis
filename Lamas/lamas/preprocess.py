import logging
import re
from dataclasses import dataclass

from openpyxl import load_workbook
from xlrd import open_workbook

from .sheet_config import Config, SheetConfig

logger = logging.getLogger(__name__)

MAGICS_RAW = [
    'סמל הרשות',
    '^מחוז',
    'מעמד מוניציפאלי',
    'מיסים ומענקים',
    'שם הרשות',
    'גירעון מצטבר בסוף שנה',
    'סמל רשות מקומית',
]
SPACES = re.compile(r'\s+')
LETTERS = re.compile('[א-ת]')
# Excel embeds invisible Unicode bidi control characters (RIGHT-TO-LEFT EMBEDDING U+202B, POP
# DIRECTIONAL FORMATTING U+202C, and similar) around numbers in RTL Hebrew sheets to force correct
# visual rendering direction - found for real, flagged as "BAD FLOATS" because float() can't
# parse a string with invisible control characters in it even though the visible digits are fine.
BIDI_CONTROL_CHARS = re.compile('[‎‏‪-‮⁦-⁩]')
MAGICS = [re.compile(x) for x in MAGICS_RAW]
MIN_SIZE = 30
HEADER_SIZE = 5


class SheetParseError(Exception):
    """Base class for sheet-parsing failures. Caught per-sheet so one bad sheet doesn't abort
    the whole multi-year run (the original code used bare `assert`, which aborted everything)."""


class MagicPositionError(SheetParseError):
    pass


class OrientationError(SheetParseError):
    pass


class NameColumnError(SheetParseError):
    pass


class BadHeaderError(SheetParseError):
    pass


class RowExtractionError(SheetParseError):
    pass


def clean(v):
    if isinstance(v, str):
        return SPACES.sub(' ', v.strip())
    return v


def get_safe(data, r, c):
    try:
        val = data[r][c]
        if val is not None:
            val = str(val).strip()
        return val
    except IndexError:
        return None


def fix_value(v, bad_floats):
    if v is not None:
        v = BIDI_CONTROL_CHARS.sub('', v).strip()
    if v in ('-', '..', '', '.', '. .', None):
        return None
    if 'http' in v:
        return None
    if LETTERS.findall(v):
        return v
    v = v.replace(',', '')
    v = v.replace('*', '')
    v = v.strip()
    if v.startswith('(') and v.endswith(')'):
        v = v[1:-1]
    if v.endswith('.0'):
        v = v[:-2]
    try:
        float(v)
    except Exception:
        bad_floats.add(v)
    return v


@dataclass
class HeaderLayout:
    get: object  # callable(r, c) -> value
    min_row: int
    min_col: int
    headers: list
    name_idx: int


def locate_headers(year, name, data, config: Config) -> HeaderLayout:
    """Find the header block for one sheet: anchor detection, orientation, multi-row header join.

    Raises a SheetParseError subclass on any layout surprise - this is the tool `lamas diagnose`
    calls directly (uncaught) to report exactly what went wrong.
    """
    magic_pos = []
    candidates = set()
    for row in range(HEADER_SIZE):
        for col, val in enumerate(data[row]):
            if isinstance(val, str):
                candidates.add(((row, col), val))
    for row, rowdata in enumerate(data):
        for col, val in enumerate(rowdata[:HEADER_SIZE]):
            if isinstance(val, str):
                candidates.add(((row, col), val))

    for magic in MAGICS:
        for pos, val in candidates:
            if magic.search(val):
                magic_pos.append([*pos, magic])
                break

    if len(magic_pos) < 2:
        raise MagicPositionError(f'Failed to find enough magic positions {year}, {name}, {magic_pos}')

    if magic_pos[0][1] == magic_pos[1][1]:
        min_row = magic_pos[0][1]
        min_col = min(p[0] for p in magic_pos)
        get = lambda r, c: get_safe(data, c, r)
    elif magic_pos[0][0] == magic_pos[1][0]:
        min_row = magic_pos[0][0]
        min_col = min(p[1] for p in magic_pos)
        get = lambda r, c: get_safe(data, r, c)
    else:
        raise OrientationError(f'invalid magic positions {year}, {name}, {magic_pos}')

    headers_size = config.header_rows
    headers = []
    headers_front = [None] * headers_size
    min_row = min_row - config.extend_headers_top

    c = 0
    name_idx = None
    while True:
        column = [get(min_row + r, min_col + c) for r in range(headers_size)]
        if not any(column):
            break
        for i, v in enumerate(column):
            if v and 'שנת עדכון' not in v:
                headers_front[i] = v
                if len(LETTERS.findall(v)) > 0:
                    for ii in range(i + 1, headers_size):
                        headers_front[ii] = None
        header = [v for v in headers_front if v]

        # Clear fully numeric headers from the front, they are not to be pulled through.
        for i, v in enumerate(headers_front):
            if v and len(LETTERS.findall(v)) == 0:
                headers_front[i] = None

        if any(f in h for h in header for f in ('שם הרשות', 'סמל הרשות')):
            headers_front = [None] * headers_size
        if any('שם הרשות' in h for h in header):
            if name_idx is not None:
                raise NameColumnError(
                    f'name_idx already set {name_idx}, {year}, {name}, {header}, {headers}'
                )
            name_idx = len(headers)
        if headers:
            try:
                float(headers[-1])
            except ValueError:
                pass
            else:
                raise BadHeaderError(f'Bad Header extracted: {year}, {name}, {header}')
        header = [x.strip() for x in header if x]
        header = [x for x in header if x and len(x) > 1]
        header = '/'.join(header)
        headers.append(header)

        c += 1
    if name_idx is None:
        raise NameColumnError(f'Failed to find name index {year}, {name}, {headers[:10]}')

    return HeaderLayout(get=get, min_row=min_row, min_col=min_col, headers=headers, name_idx=name_idx)


def process_sheet(year, name, data, filename, sheet_config: SheetConfig):
    name = name.strip()
    config = sheet_config.get(year, name)
    if config.skip:
        return
    logger.info('PROCESSING %r %r %r', year, name, config)

    layout = locate_headers(year, name, data, config)
    get = layout.get
    min_row = layout.min_row
    min_col = layout.min_col
    headers = layout.headers
    name_idx = layout.name_idx
    headers_size = config.header_rows

    # extend_headers_bottom skips N rows immediately after the header block before real
    # per-municipality data starts (e.g. a blank spacer row followed by nationwide/aggregate
    # summary rows inserted above the first real row - a layout CBS introduced in 2023).
    data_start = min_row + headers_size + config.extend_headers_bottom

    r = data_start
    max_r = None
    while True:
        row = [get(r, min_col + c) for c in range(len(headers))]
        if not any(row):
            break
        if max_r is None:
            max_r = r
        if len([x for x in row if x]) > MIN_SIZE / 2:
            max_r = r
        r += 1

    r = data_start
    bad_floats = set()
    while max_r is not None and r <= max_r:
        row = [get(r, min_col + c) for c in range(len(headers))]
        try:
            muni_name = row[name_idx].strip().replace('*', '')
        except Exception as e:
            raise RowExtractionError(
                f'Failed to read municipality name at row {r}: {config}, {year}, {name!r}, {row}'
            ) from e
        for h, v in zip(headers, row):
            if h != headers[name_idx]:
                v = fix_value(v, bad_floats)
                yield dict(
                    year=year,
                    sheet=name,
                    name=muni_name,
                    header=h,
                    value=v,
                    filename=filename,
                )
        r += 1
    if len(bad_floats) > 0:
        logger.warning('BAD FLOATS %s', bad_floats)


def _process_sheet_isolated(year, name, data, filename, sheet_config):
    try:
        yield from process_sheet(year, name, data, filename, sheet_config)
    except SheetParseError as e:
        logger.error('Skipping sheet %r/%r after parse failure: %s', year, name, e)


def preprocess_file(year, filename, sheet_config: SheetConfig):
    if filename.endswith('.xlsx'):
        wb = load_workbook(filename)
        for name in wb.sheetnames:
            sheet = wb[name]
            if sheet.max_column < MIN_SIZE or sheet.max_row < MIN_SIZE:
                continue
            data = [[clean(cell.value) for cell in row] for row in sheet.rows]
            yield from _process_sheet_isolated(year, sheet.title, data, filename, sheet_config)
    else:
        wb = open_workbook(filename)
        for idx in range(wb.nsheets):
            sheet = wb.sheet_by_index(idx)
            if sheet.nrows < MIN_SIZE or sheet.ncols < MIN_SIZE:
                continue
            data = [[clean(sheet.cell_value(r, c)) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
            yield from _process_sheet_isolated(year, sheet.name, data, filename, sheet_config)


def preprocess_files(filenames, sheet_config: SheetConfig = None):
    sheet_config = sheet_config or SheetConfig.load()
    for year, filename in filenames.items():
        logger.info('%s %s', year, filename)
        yield from preprocess_file(year, filename, sheet_config)
