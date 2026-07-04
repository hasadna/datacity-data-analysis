import logging

from openpyxl import load_workbook
from xlrd import open_workbook

from .preprocess import MIN_SIZE, SheetParseError, clean, locate_headers, process_sheet
from .sheet_config import SheetConfig

logger = logging.getLogger(__name__)


def _load_sheets(filename):
    if filename.endswith('.xlsx'):
        wb = load_workbook(filename)
        for name in wb.sheetnames:
            sheet = wb[name]
            if sheet.max_column < MIN_SIZE or sheet.max_row < MIN_SIZE:
                continue
            data = [[clean(cell.value) for cell in row] for row in sheet.rows]
            yield sheet.title, data
    else:
        wb = open_workbook(filename)
        for idx in range(wb.nsheets):
            sheet = wb.sheet_by_index(idx)
            if sheet.nrows < MIN_SIZE or sheet.ncols < MIN_SIZE:
                continue
            data = [[clean(sheet.cell_value(r, c)) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
            yield sheet.name, data


def diagnose_sheet(year, name, data, sheet_config: SheetConfig, preview_size=15):
    """Dry-run parse of one sheet against the current sheet_config.yaml. Never raises - any
    SheetParseError is captured in the result so a human/skill can decide how to tune config.

    Runs the FULL process_sheet() (header location + row extraction), not just header location -
    a sheet can locate headers fine and still fail later on a row that isn't a real municipality
    (e.g. a "Total" summary row), which only shows up once you actually walk the data rows.
    """
    stripped = name.strip()
    config = sheet_config.get(year, stripped)
    result = {
        'year': year,
        'sheet': name,
        'config': config,
        'skipped': config.skip,
        'error': None,
        'name_idx': None,
        'header_count': None,
        'header_preview': [],
        'data_row_count': None,
        'row_count': len(data),
        'col_count': len(data[0]) if data else 0,
    }
    if config.skip:
        return result
    try:
        layout = locate_headers(year, stripped, data, config)
    except SheetParseError as e:
        result['error'] = f'{type(e).__name__}: {e}'
        return result
    result['name_idx'] = layout.name_idx
    result['header_count'] = len(layout.headers)
    seen = []
    for h in layout.headers:
        if h and h not in seen:
            seen.append(h)
        if len(seen) >= preview_size:
            break
    result['header_preview'] = seen

    try:
        extracted = list(process_sheet(year, name, data, filename='<diagnose>', sheet_config=sheet_config))
    except SheetParseError as e:
        result['error'] = f'{type(e).__name__}: {e}'
        return result
    result['data_row_count'] = len({(r['name']) for r in extracted})
    return result


def diagnose_year(year, filename, sheet_config: SheetConfig, sheet_filter=None):
    results = []
    for name, data in _load_sheets(filename):
        if sheet_filter and name.strip() != sheet_filter.strip():
            continue
        results.append(diagnose_sheet(year, name, data, sheet_config))
    return results
